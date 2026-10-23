"""Excel 数据集：模板、解析、导出。只服务普通对话型测试集（一行一轮）。

分组规则：同名「用例名称」的行按「轮次」升序合成一个用例；跨用例的轮次号互不影响。
断言合成：期望包含→answer_contains（换行分隔多个）；期望不包含→answer_not_contains；
期望工具→tool_called；禁止工具→tool_forbidden；四者全空 = 缺少必过规则（行错误）。

断言字段名以 ``schemas.py`` 实际定义为准：tool_called 是 ``name``（Identifier）加必填
``min_calls``，tool_forbidden 是 ``names`` 列表；不是计划草稿里的单一 ``tool`` 字段。
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook, load_workbook

COLUMNS = ("用例名称", "轮次", "输入", "期望包含", "期望不包含", "期望工具", "禁止工具", "标签")
MAX_ROWS = 500
MAX_BYTES = 5 * 1024 * 1024


def build_template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    if ws is None:  # pragma: no cover - 新 Workbook 必有默认 sheet
        ws = wb.create_sheet()
    ws.title = "用例"
    ws.append(COLUMNS)
    ws.append(["示例用例", 1, "你好，请用一句话介绍知序。", "知序", "", "", "", "示例"])
    info = wb.create_sheet("说明")
    info["A1"] = "数据集名称"
    info["B1"] = "我的数据集"
    info["A3"] = "一行一轮；同名多行按轮次合成多轮用例。期望包含可写多行（换行分隔）。"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(row: tuple[Any, ...], index: int) -> str:
    value = row[index] if index < len(row) else None
    return "" if value is None else str(value).strip()


def _split_lines(value: str) -> list[str]:
    return [line.strip() for line in value.splitlines() if line.strip()]


def _turn_no(row: tuple[Any, ...]) -> int | None:
    """轮次列：空按 0（单行用例的常见写法）；非数字返回 None，由调用方记行错误。"""
    text = _cell(row, 1)
    if not text:
        return 0
    try:
        return int(float(text))
    except ValueError:
        return None


def parse_workbook(data: bytes) -> dict[str, Any]:
    if len(data) > MAX_BYTES:
        return {"suite": None, "cases": [], "errors": [{"row": 0, "message": "文件超过 5MB"}]}
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        return {"suite": None, "cases": [], "errors": [{"row": 0, "message": f"无法解析 xlsx: {exc}"}]}
    ws = wb["用例"] if "用例" in wb.sheetnames else wb.active
    if ws is None:
        return {"suite": None, "cases": [], "errors": [{"row": 0, "message": "没有可读取的工作表"}]}
    suite_name = "导入数据集"
    if "说明" in wb.sheetnames and wb["说明"]["B1"].value:
        suite_name = str(wb["说明"]["B1"].value).strip() or suite_name

    grouped: dict[str, list[tuple[int, int, tuple[Any, ...]]]] = {}
    errors: list[dict[str, Any]] = []
    for row_index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(_cell(row, i) == "" for i in range(len(COLUMNS))):
            continue
        name = _cell(row, 0)
        if not name:
            errors.append({"row": row_index, "message": "缺少用例名称"})
            continue
        turn_no = _turn_no(row)
        if turn_no is None:
            errors.append({"row": row_index, "message": "轮次需为数字"})
            continue
        grouped.setdefault(name, []).append((row_index, turn_no, row))

    cases: list[dict[str, Any]] = []
    for name, entries in grouped.items():
        entries.sort(key=lambda item: (item[1], item[0]))
        turns: list[dict[str, str]] = []
        contains: list[str] = []
        not_contains: list[str] = []
        tools: list[str] = []
        forbidden: list[str] = []
        tags: list[str] = []
        for row_index, _turn, row in entries:
            content = _cell(row, 2)
            if not content:
                errors.append({"row": row_index, "message": "缺少输入"})
                continue
            turns.append({"content": content})
            contains += _split_lines(_cell(row, 3))
            not_contains += _split_lines(_cell(row, 4))
            tools += [t for t in _cell(row, 5).replace("，", ",").split(",") if t]
            forbidden += [t for t in _cell(row, 6).replace("，", ",").split(",") if t]
            tags += [t for t in _cell(row, 7).replace("，", ",").split(",") if t]
        if not turns:
            continue
        assertions: list[dict[str, Any]] = []
        for index, text in enumerate(dict.fromkeys(contains)):
            assertions.append({"id": f"c{index}", "kind": "answer_contains", "required": True,
                               "text": text, "case_sensitive": False})
        for index, text in enumerate(dict.fromkeys(not_contains)):
            assertions.append({"id": f"n{index}", "kind": "answer_not_contains", "required": True,
                               "text": text, "case_sensitive": False})
        # schemas.py：tool_called 字段是 name + 必填 min_calls（导入语义取 1 次）；tool_forbidden 是 names 列表。
        for index, tool in enumerate(dict.fromkeys(tools)):
            assertions.append({"id": f"t{index}", "kind": "tool_called", "required": True,
                               "name": tool, "min_calls": 1})
        for index, tool in enumerate(dict.fromkeys(forbidden)):
            assertions.append({"id": f"f{index}", "kind": "tool_forbidden", "required": True,
                               "names": [tool]})
        if not assertions:
            errors.append({"row": entries[0][0], "message": f"用例 {name} 缺少期望包含/期望工具等必过规则"})
            continue
        cases.append({
            "name": name[:120], "description": "", "tags": tags[:20], "enabled": True,
            "turns": turns[:20], "fixtures": [], "memory_seed": [], "clock_iso": None,
            "tool_fixtures": [],
            "execution": {"tool_mode": "live",
                          "allowed_tools": ["read_file", "write_file", "edit_file", "list_dir",
                                            "web_search", "web_fetch"],
                          "max_iterations": 12, "timeout_s": 120},
            "assertions": assertions[:50], "rubric": None, "script": [], "artifact_paths": [],
        })
    suite_payload = {"name": suite_name[:120], "description": "Excel 导入",
                     "tags": ["dataset"], "base_revision": None}
    return {"suite": suite_payload, "cases": cases[:200], "errors": errors}


def export_suite_workbook(suite: dict[str, Any], cases: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    if ws is None:  # pragma: no cover - 新 Workbook 必有默认 sheet
        ws = wb.create_sheet()
    ws.title = "用例"
    ws.append(COLUMNS)
    for case in cases:
        turns = case.get("turns") or []
        assertions = case.get("assertions") or []
        contains = [str(a.get("text", "")) for a in assertions
                    if a.get("kind") == "answer_contains"]
        not_contains = [str(a.get("text", "")) for a in assertions
                        if a.get("kind") == "answer_not_contains"]
        tools = [str(a.get("name", "")) for a in assertions if a.get("kind") == "tool_called"]
        forbidden: list[str] = []
        for assertion in assertions:
            if assertion.get("kind") == "tool_forbidden":
                forbidden.extend(str(n) for n in (assertion.get("names") or []))
        tags = ",".join(str(tag) for tag in (case.get("tags") or []))
        for index, turn in enumerate(turns, start=1):
            ws.append([
                str(case.get("name", "")), index, str(turn.get("content", "")),
                "\n".join(contains) if index == 1 else "",
                "\n".join(not_contains) if index == 1 else "",
                ",".join(tools) if index == 1 else "",
                ",".join(forbidden) if index == 1 else "",
                tags if index == 1 else "",
            ])
    info = wb.create_sheet("说明")
    info["A1"] = "数据集名称"
    info["B1"] = str(suite.get("name", ""))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
