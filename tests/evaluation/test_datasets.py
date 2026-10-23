"""Excel 数据集纯函数测试：模板 / 解析 / 导出（nanobot/evaluation/datasets.py）。"""

from __future__ import annotations

import io
import uuid

from openpyxl import Workbook, load_workbook

from nanobot.evaluation.datasets import (
    COLUMNS,
    MAX_BYTES,
    build_template_bytes,
    export_suite_workbook,
    parse_workbook,
)
from nanobot.evaluation.schemas import CaseWrite


def _fresh_sheet() -> "tuple[Workbook, object]":
    """只含表头「用例」sheet 的最小工作簿，供逐条测试填行。"""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "用例"
    ws.append(COLUMNS)
    return wb, ws


def test_template_roundtrip() -> None:
    parsed = parse_workbook(build_template_bytes())
    assert parsed["errors"] == []  # 模板自带示例行合法
    assert parsed["cases"][0]["name"] == "示例用例"
    assert parsed["suite"]["name"] == "我的数据集"  # 说明页 B1 → suite 名称


def test_rows_group_into_multiturn_case() -> None:
    data = build_template_bytes()
    # 用 openpyxl 改写：同名两行、轮次 1/2 → 一个用例 2 turns
    wb = load_workbook(io.BytesIO(data))
    ws = wb["用例"]
    ws.append(["多轮用例", 1, "第一问", "答案一", "", "", "", ""])
    ws.append(["多轮用例", 2, "第二问", "答案二", "", "", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    parsed = parse_workbook(buf.getvalue())
    assert parsed["errors"] == []
    case = next(c for c in parsed["cases"] if c["name"] == "多轮用例")
    assert len(case["turns"]) == 2
    assert [t["content"] for t in case["turns"]] == ["第一问", "第二问"]


def test_missing_required_rule_is_row_error() -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "用例"
    ws.append(COLUMNS)
    ws.append(["空规则", 1, "hi", "", "", "", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    parsed = parse_workbook(buf.getvalue())
    assert parsed["cases"] == []
    assert any("缺少" in e["message"] for e in parsed["errors"])


def test_multiline_contains_split_into_assertions() -> None:
    wb, ws = _fresh_sheet()
    ws.append(["多行包含", 1, "hi", "alpha\nbeta\n", "", "", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    parsed = parse_workbook(buf.getvalue())
    assertions = parsed["cases"][0]["assertions"]
    contains = [a["text"] for a in assertions if a["kind"] == "answer_contains"]
    assert contains == ["alpha", "beta"]


def test_non_numeric_turn_is_row_error() -> None:
    wb, ws = _fresh_sheet()
    ws.append(["坏轮次", "第一轮", "hi", "ok", "", "", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    parsed = parse_workbook(buf.getvalue())
    assert parsed["cases"] == []
    assert any("轮次" in e["message"] for e in parsed["errors"])


def test_empty_and_garbage_bytes_return_row_zero_errors() -> None:
    empty = parse_workbook(b"")
    assert empty["suite"] is None
    assert empty["cases"] == []
    assert empty["errors"] and empty["errors"][0]["row"] == 0
    garbage = parse_workbook(b"this is not an xlsx")
    assert garbage["suite"] is None
    assert garbage["cases"] == []
    assert any("无法解析" in e["message"] for e in garbage["errors"])


def test_oversized_file_rejected_with_5mb_semantics() -> None:
    # 路由层把它映射为 422；纯函数契约是 row=0 的结构化错误，不抛异常。
    parsed = parse_workbook(b"\x00" * (MAX_BYTES + 1))
    assert parsed["suite"] is None
    assert parsed["cases"] == []
    assert "5MB" in parsed["errors"][0]["message"]
    assert parsed["errors"][0]["row"] == 0


def test_parsed_case_passes_casewrite_schema() -> None:
    """字段名以 schemas 为准：tool_called.name/min_calls、tool_forbidden.names 必须能过 CaseWrite。"""
    wb, ws = _fresh_sheet()
    ws.append(["schema用例", 1, "第一问", "知序", "", "web_search", "exec", "标签A"])
    ws.append(["schema用例", 2, "第二问", "", "机密", "", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    parsed = parse_workbook(buf.getvalue())
    assert parsed["errors"] == []

    case = parsed["cases"][0]
    case_write = CaseWrite.model_validate({**case, "suite_id": uuid.uuid4()})
    kinds = {a.kind for a in case_write.assertions}
    assert kinds == {"answer_contains", "answer_not_contains", "tool_called", "tool_forbidden"}
    tool_assertion = next(a for a in case_write.assertions if a.kind == "tool_called")
    assert tool_assertion.name == "web_search"
    assert tool_assertion.min_calls == 1
    forbidden = next(a for a in case_write.assertions if a.kind == "tool_forbidden")
    assert forbidden.names == ["exec"]


def test_export_roundtrip() -> None:
    parsed = parse_workbook(build_template_bytes())
    exported = export_suite_workbook(parsed["suite"], parsed["cases"])
    reparsed = parse_workbook(exported)
    assert reparsed["errors"] == []
    assert reparsed["suite"]["name"] == "我的数据集"
    case = reparsed["cases"][0]
    assert case["name"] == "示例用例"
    assert [t["content"] for t in case["turns"]] == ["你好，请用一句话介绍知序。"]
    assert case["tags"] == ["示例"]
    contains = [a["text"] for a in case["assertions"] if a["kind"] == "answer_contains"]
    assert contains == ["知序"]
